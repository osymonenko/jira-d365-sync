"""Regression test: inserting rows into an EARLIER week can force overflow rows
to be inserted right before a LATER week's block, physically shifting that
block (including an already-written, correct check-sum row) down. openpyxl's
insert_rows() moves cell values/formulas as-is without rewriting formula text,
so the shifted check-sum row's SUM() formula keeps its old (now wrong)
absolute row references unless something re-derives it after all insertions
for the run are done. This is what _reconcile_check_sums() exists to fix."""

import importlib.util
import pathlib
import tempfile
from datetime import datetime

import openpyxl

_p = pathlib.Path(__file__).parent / "jira-sync.py"
_spec = importlib.util.spec_from_file_location("jira_sync", _p)
jira_sync = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(jira_sync)


def build_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1).value = "weekstart"
    ws.cell(1, 3).value = "Task"
    # Block A: date row 2, no tasks yet, zero blank rows before Block B — tight
    # enough that adding tasks to A forces an overflow insert before B.
    ws.cell(2, 1).value = datetime(2026, 7, 5)
    ws.cell(2, 2).value = datetime(2026, 7, 11)
    # Block B: date row 3, one pre-existing task at row 4.
    ws.cell(3, 1).value = datetime(2026, 7, 12)
    ws.cell(3, 2).value = datetime(2026, 7, 18)
    ws.cell(4, 3).value = "Existing Task"
    wb.save(path)


def main():
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    path = tmp.name
    build_workbook(path)

    insertions = {
        "2026-07-12": [  # Block B — processed FIRST (reversed/bottom-up order)
            {"name": "New Task", "hours_by_col": {4: 1.0}},
        ],
        "2026-07-05": [  # Block A — processed SECOND, its overflow shifts B down
            {"name": "Task 1", "hours_by_col": {4: 0.5}},
            {"name": "Task 2", "hours_by_col": {4: 0.5}},
            {"name": "Task 3", "hours_by_col": {4: 0.5}},
        ],
    }
    total = jira_sync.insert_standard_rows(path, insertions)
    assert total == 4, f"expected 4 rows written, got {total}"

    wb = openpyxl.load_workbook(path)  # formulas as text, not evaluated
    ws = wb.worksheets[0]

    def c(r, col):
        return ws.cell(r, col).value

    weeks = jira_sync.find_weeks(ws)
    week_a = next(w for w in weeks if w["week_start"] == datetime(2026, 7, 5).date())
    week_b = next(w for w in weeks if w["week_start"] == datetime(2026, 7, 12).date())

    # Block A: fresh check sum, unaffected by any of this — sanity check only.
    a_check = week_a["check_sum_row"]
    assert c(a_check, 3) == "check sum", c(a_check, 3)
    assert c(a_check, 4) == f"=SUM(D{week_a['start_row']}:D{a_check - 1})", c(a_check, 4)
    print(f"  OK block A: check sum at row {a_check} references its own block")

    # Block B: shifted down by A's overflow. Its check-sum formula MUST reference
    # its own (post-shift) task rows, not the stale pre-shift row numbers.
    b_check = week_b["check_sum_row"]
    assert b_check is not None, "block B lost its check-sum row"
    task_rows = [
        r for r in range(week_b["start_row"], week_b["end_row"])
        if r != b_check and c(r, 3) not in (None, "", "check sum")
    ]
    expected_first, expected_last = min(task_rows), max(task_rows)
    actual_formula = c(b_check, 4)
    expected_formula = f"=SUM(D{expected_first}:D{expected_last})"
    assert actual_formula == expected_formula, (
        f"block B check sum at row {b_check} references stale rows: "
        f"got {actual_formula!r}, expected {expected_formula!r} "
        f"(task rows are actually {task_rows})"
    )
    print(f"  OK block B: check sum at row {b_check} correctly references "
          f"its post-shift task rows {task_rows}, not stale pre-shift rows")

    wb.close()
    print("ALL PASS")


if __name__ == "__main__":
    main()

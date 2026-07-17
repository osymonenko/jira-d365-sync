"""Network-free test for standard-row insertion (hours land in day columns)."""

import importlib.util
import pathlib
import tempfile
from datetime import datetime

import openpyxl

_p = pathlib.Path(__file__).parent / "jira-sync.py"
_spec = importlib.util.spec_from_file_location("jira_sync", _p)
jira_sync = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(jira_sync)


def build_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1).value = "weekstart"
    ws.cell(1, 3).value = "Task"
    ws.cell(2, 1).value = datetime(2026, 7, 12)   # Sunday
    ws.cell(2, 2).value = datetime(2026, 7, 18)
    ws.cell(2, 3).value = "Internal Daily meeting"  # pre-existing (dedup, diff case)
    ws.cell(3, 3).value = "check sum"
    wb.save(path)


def main():
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    path = tmp.name
    build_workbook(path)

    insertions = {
        "2026-07-12": [
            {"name": "internal daily meeting", "hours_by_col": {4: 0.5}},  # dup (case-insensitive) -> skipped
            {"name": "Weekly project report", "hours_by_col": {8: 1.0}},
            {"name": "Bug verification", "hours_by_col": {}},              # QA placeholder, no hours
        ]
    }
    total = jira_sync.insert_standard_rows(path, insertions)
    assert total == 2, f"expected 2 inserted (dup skipped), got {total}"
    print("  OK dedup skipped the case-different duplicate")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    # Collect name -> {col: value} for column C rows and their day columns.
    found = {}
    for r in range(1, ws.max_row + 1):
        name = ws.cell(r, 3).value
        if not name:
            continue
        found[name] = {c: ws.cell(r, c).value for c in (4, 5, 6, 7, 8) if ws.cell(r, c).value is not None}
    wb.close()

    assert found.get("Weekly project report") == {8: 1}, found.get("Weekly project report")
    print("  OK weekly report hours landed in Friday column (8)")
    assert "Bug verification" in found and found["Bug verification"] == {}, found.get("Bug verification")
    print("  OK QA placeholder inserted with no hours")

    print("ALL PASS")


if __name__ == "__main__":
    main()

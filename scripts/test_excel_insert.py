"""Network-free test for update_excel row insertion, incl. weeks without a
'check sum' row (regression: such weeks used to be silently skipped)."""

import importlib.util
import pathlib
import tempfile
from datetime import datetime

import openpyxl

# Load the hyphenated module by path (scripts/ is sys.path[0] when run as
# `python scripts/test_excel_insert.py`, so its `from month_filter import` works).
_p = pathlib.Path(__file__).parent / "jira-sync.py"
_spec = importlib.util.spec_from_file_location("jira_sync", _p)
jira_sync = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(jira_sync)


def build_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    # Row 1: header (non-date col A -> not a week block)
    ws.cell(1, 1).value = "weekstart"
    ws.cell(1, 3).value = "Task"
    # Row 2-3: week A WITH a check sum row
    ws.cell(2, 1).value = datetime(2026, 6, 28)
    ws.cell(2, 2).value = datetime(2026, 7, 4)
    ws.cell(2, 3).value = "Meeting A"
    ws.cell(3, 3).value = "check sum"
    # Row 4-5: week B with tasks but NO check sum (blank row 5)
    ws.cell(4, 1).value = datetime(2026, 7, 5)
    ws.cell(4, 2).value = datetime(2026, 7, 11)
    ws.cell(4, 3).value = "Meeting B"
    # Row 6: week C — bare date row only, no check sum
    ws.cell(6, 1).value = datetime(2026, 7, 12)
    ws.cell(6, 2).value = datetime(2026, 7, 18)
    wb.save(path)


def names_by_week(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    weeks = jira_sync.find_weeks(wb.worksheets[0])
    wb.close()
    return {jira_sync.fmt(w["week_start"]): w for w in weeks}


def main():
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    path = tmp.name
    build_workbook(path)

    url = "https://example/x"
    insertions = {
        "2026-06-28": [("New A", url)],
        "2026-07-05": [("New B", url)],
        "2026-07-12": [("New C", url)],
    }
    total = jira_sync.update_excel(path, insertions)

    assert total == 3, f"expected 3 rows inserted, got {total}"
    print("  OK total == 3")

    weeks = names_by_week(path)

    assert "New A" in weeks["2026-06-28"]["existing_names"], "New A missing from week A"
    print("  OK week A (has check sum) got its row")

    assert "New B" in weeks["2026-07-05"]["existing_names"], "New B missing from week B"
    print("  OK week B (no check sum) got its row")

    assert "New C" in weeks["2026-07-12"]["existing_names"], "New C missing from week C"
    print("  OK week C (bare date, no check sum) got its row")

    # Week A's check sum row must still be intact after insertion.
    assert weeks["2026-06-28"]["check_sum_row"] is not None, "week A lost its check sum row"
    print("  OK week A check sum row preserved")

    # Rows must land in the correct block, not bleed into a neighbour.
    assert "New B" not in weeks["2026-06-28"]["existing_names"], "New B bled into week A"
    assert "New C" not in weeks["2026-07-05"]["existing_names"], "New C bled into week B"
    print("  OK no cross-block bleed")

    test_hyperlinks_survive_multiweek_insert()

    print("ALL PASS")


def build_hyperlink_workbook(path):
    """Two weeks, each with a check sum and one PRE-EXISTING hyperlinked row
    (as if inserted by an earlier run)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1).value = "weekstart"
    # Week A
    ws.cell(2, 1).value = datetime(2026, 6, 28)
    ws.cell(2, 2).value = datetime(2026, 7, 4)
    ws.cell(2, 3).value = "A meeting"
    a_old = ws.cell(3, 3)
    a_old.value = "A existing"
    a_old.hyperlink = "http://existing/A"
    ws.cell(4, 3).value = "check sum"
    # Week B
    ws.cell(5, 1).value = datetime(2026, 7, 5)
    ws.cell(5, 2).value = datetime(2026, 7, 11)
    ws.cell(5, 3).value = "B meeting"
    b_old = ws.cell(6, 3)
    b_old.value = "B existing"
    b_old.hyperlink = "http://existing/B"
    ws.cell(7, 3).value = "check sum"
    wb.save(path)


def test_hyperlinks_survive_multiweek_insert():
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    path = tmp.name
    build_hyperlink_workbook(path)

    # Insert into BOTH weeks in one run — the upper-week insert shifts the
    # lower week's rows and used to orphan its hyperlinks.
    insertions = {
        "2026-06-28": [("A new", "http://new/A")],
        "2026-07-05": [("B new", "http://new/B")],
    }
    jira_sync.update_excel(path, insertions)

    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    # Map value -> hyperlink target for every cell in column C.
    links = {}
    orphans = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink is None:
                continue
            if cell.value is None:
                orphans.append((cell.coordinate, cell.hyperlink.target))
            else:
                links[cell.value] = cell.hyperlink.target
    wb.close()

    expected = {
        "A existing": "http://existing/A",
        "B existing": "http://existing/B",
        "A new": "http://new/A",
        "B new": "http://new/B",
    }
    for name, target in expected.items():
        assert links.get(name) == target, \
            f"{name!r}: expected link {target}, got {links.get(name)!r}"
    print("  OK all 4 rows keep the correct hyperlink after multi-week insert")

    assert not orphans, f"orphaned hyperlinks on empty cells: {orphans}"
    print("  OK no orphaned hyperlinks left on blank rows")


if __name__ == "__main__":
    main()
